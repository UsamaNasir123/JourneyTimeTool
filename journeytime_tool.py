#!/usr/bin/env python3
"""
Journey-time monthly file builder  (core + command line)

Input : Google travel-time CSV  (columns: eui, date_created, duration_in_traffic)
Output: one Excel file per month  result_<eui>_<date>_<date>.xlsx  (+ estimate_log.csv)
        - ONE day per month, preferred weekday (Thursday) first
        - 06:00 .. 20:50, 10-minute rows, layout of sample_journeytime_before_after.xlsx
        - missing readings are PREDICTED from the typical Thursday pattern of your own data
        - months with no data at all get a predicted Thursday (optional)
        - every predicted / non-measured cell is highlighted yellow (optional) and listed in the log

CLI:  python journeytime_tool.py INPUT.csv -o OUT_DIR [--weekday Thu] [--min-coverage 0.2]
                                 [--before BEFORE_FILE] [--no-fill-months] [--no-highlight] [--force-weekday]
"""
import argparse, io, math, os, re, sys, zipfile, zlib
from datetime import date, datetime, time, timedelta
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

STEP = 10
SLOT_MIN = list(range(6 * 60, 21 * 60, STEP))              # 06:00 .. 20:50 -> 90 rows
SLOTS = [time(m // 60, m % 60) for m in SLOT_MIN]
WD = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
YELLOW = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")


# ---------------------------------------------------------------- loading
def load(csv_or_df, value="duration_in_traffic"):
    df = csv_or_df if isinstance(csv_or_df, pd.DataFrame) else pd.read_csv(csv_or_df, encoding="utf-8-sig")
    for c in ("eui", "date_created", value):
        if c not in df.columns:
            raise ValueError(f"CSV must contain column '{c}'. Found: {list(df.columns)}")
    df = df.copy()
    df["dt"] = pd.to_datetime(df["date_created"], errors="coerce")
    df["v"] = pd.to_numeric(df[value], errors="coerce")
    df = df.dropna(subset=["dt", "v"])
    df = df[(df.dt.dt.hour >= 6) & (df.dt.dt.hour < 21)]
    df["date"] = df.dt.dt.date
    return df


def measured_day(g):
    """90-slot Series: each reading goes to the nearest 10-min slot (NaN where nothing was recorded)."""
    day = pd.Timestamp(g["date"].iloc[0])
    idx = pd.DatetimeIndex([day + pd.Timedelta(minutes=m) for m in SLOT_MIN])
    pos = idx.get_indexer(pd.DatetimeIndex(g["dt"]), method="nearest", tolerance=pd.Timedelta(minutes=STEP / 2))
    ok = pos >= 0
    out = pd.Series(np.nan, index=range(len(SLOTS)))
    if ok.any():
        m = pd.Series(g["v"].values[ok]).groupby(pos[ok]).mean()
        out.loc[m.index] = m.values
    return out


def load_before(src, before_date=None):
    """Read the fixed 'Before' day (csv or xlsx: a Time column + a '(Before)' column, any times).
    Returns (date, 90-slot Series with NaN where the file has no reading)."""
    if hasattr(src, "name") and str(src.name).lower().endswith((".xlsx", ".xls")) or str(src).lower().endswith((".xlsx", ".xls")):
        b = pd.read_excel(src)
    else:
        b = pd.read_csv(src, encoding="utf-8-sig")
    tcol = next((c for c in b.columns if str(c).strip().lower() == "time"), b.columns[0])
    vcols = [c for c in b.columns if c != tcol]
    bcol = next((c for c in vcols if "before" in str(c).lower()), vcols[0] if vcols else None)
    if bcol is None:
        raise ValueError("Before file needs a 'Time' column and a '(Before)' column.")
    if before_date is None:
        m = re.search(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", str(bcol))
        if not m:
            raise ValueError(f"Could not find a date like 11-09-2024 in the column name '{bcol}'. Enter the Before date manually.")
        before_date = date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    out = pd.Series(np.nan, index=range(len(SLOTS)))
    for t, v in zip(b[tcol], pd.to_numeric(b[bcol], errors="coerce")):
        if pd.isna(v) or pd.isna(t):
            continue
        t = t if isinstance(t, (time, datetime)) else pd.to_datetime(str(t)).time()
        mins = t.hour * 60 + t.minute + (1 if t.second >= 30 else 0)
        i = round((mins - SLOT_MIN[0]) / STEP)
        if 0 <= i < len(SLOTS):
            out.iloc[i] = v
    return before_date, out


# ---------------------------------------------------------------- prediction
def typical_profile(days, weekday):
    """Median value per slot over (nearly) complete days; prefers the chosen weekday."""
    full = {d: s for d, (s, _) in days.items() if s.notna().mean() >= 0.95}
    wd = [s for d, s in full.items() if WD[d.weekday()] == weekday]
    pool = wd if len(wd) >= 3 else list(full.values())
    if not pool:
        return None, "no complete day in file"
    prof = pd.concat(pool, axis=1).median(axis=1)
    basis = f"{len(pool)} complete {weekday if pool is wd else 'any-weekday'} day(s)"
    return prof.rolling(3, center=True, min_periods=1).mean(), basis   # light smoothing


def complete_day(meas, prof, decay=6, interior="profile"):
    """Fill missing slots: typical profile x (this day's own level), blended smoothly at the edges."""
    obs = meas.notna()
    if prof is None or obs.sum() < 5:
        return meas.interpolate(limit_area=None).bfill().ffill()
    r = meas / prof
    scale = r[obs].median()
    r_in = r.interpolate(limit_area="inside")
    first, last = obs.idxmax(), obs[::-1].idxmax()
    for i in range(len(r_in)):
        if i < first:
            r_in[i] = scale + (r[first] - scale) * math.exp(-(first - i) / decay)
        elif i > last:
            r_in[i] = scale + (r[last] - scale) * math.exp(-(i - last) / decay)
    est = prof * r_in
    if interior == "linear":          # between real readings just bridge them smoothly (used for the Before day)
        inside = meas.interpolate(limit_area="inside")
        est = est.where((est.index < first) | (est.index > last), inside)
    return meas.where(obs, est)


def estimate_rel_sd(series):
    """Typical slot-to-slot scatter of REAL readings, as a share of the level (None if too few pairs)."""
    x = series.values.astype(float)
    pair = ~np.isnan(x[1:]) & ~np.isnan(x[:-1])
    if pair.sum() < 8:
        return None
    d = (x[1:] - x[:-1])[pair]
    return float(np.std(d) / math.sqrt(2) / np.nanmean(x))


def add_variation(vals, real_mask, rel_sd, seed):
    """Give PREDICTED slots the natural scatter real readings have (smooth estimates look artificial).
    Real readings are never changed. Same seed -> same numbers on every run."""
    rng = np.random.default_rng(zlib.crc32(seed.encode()))
    z = np.clip(rng.standard_normal(len(vals)), -2.5, 2.5)
    out = vals.copy().astype(float)
    pred = ~np.asarray(real_mask, dtype=bool)
    out[pred] = out[pred] * (1 + rel_sd * z[pred])
    return out


# ---------------------------------------------------------------- main processing
def first_weekday(y, m, wd_idx):
    d = date(y, m, 1)
    return d + timedelta(days=(wd_idx - d.weekday()) % 7)


def process(df, weekday="Thu", min_coverage=0.2, fill_missing_months=True, force_weekday=False, before=None, natural=True):
    """Returns list of dicts: {name, day, eui, values(list of 90 floats), measured(list of bool), source, month}."""
    weekday = weekday[:3].title()
    wd_idx = WD.index(weekday)
    results = []
    b_date, b_meas = before if before is not None else (None, None)
    for eui, g_eui in df.groupby("eui"):
        days = {}
        for d, g in g_eui.groupby("date"):
            s = measured_day(g)
            days[d] = (s, s.notna().mean())
        if not days:
            continue
        prof, basis = typical_profile(days, weekday)
        rels = [r for r in (estimate_rel_sd(sr) for sr, _ in days.values() if sr.notna().mean() >= 0.95) if r]
        rel_g = float(np.median(rels)) if rels else 0.03
        ms = sorted({(d.year, d.month) for d in days})
        y, m = ms[0]
        all_months = []
        while (y, m) <= ms[-1]:
            all_months.append((y, m)); y, m = (y + 1, 1) if m == 12 else (y, m + 1)

        if b_meas is not None:
            b_vals = complete_day(b_meas, prof, interior="linear")
            if natural:
                b_vals = add_variation(b_vals, b_meas.notna(), estimate_rel_sd(b_meas) or rel_g, f"{eui}|before|{b_date}")
            b_vals = [None if pd.isna(v) else int(round(v)) for v in b_vals]
            b_mask = [bool(x) for x in b_meas.notna()]
        for (y, m) in all_months:
            cand = [d for d in days if (d.year, d.month) == (y, m) and days[d][1] >= min_coverage]
            pref = [d for d in cand if d.weekday() == wd_idx]
            if cand and (pref or not force_weekday):
                pool = pref or cand
                d = sorted(pool, key=lambda x: (-days[x][1], x))[0]
                meas, cov = days[d]
                vals = complete_day(meas, prof)
                if natural:
                    vals = add_variation(vals, meas.notna(), estimate_rel_sd(meas) or rel_g, f"{eui}|after|{d}")
                src = "measured" if cov >= 0.999 else f"measured {cov:.0%} + predicted"
                if not pref:
                    src += f" (no {weekday} data, used {WD[d.weekday()]})"
                mask = meas.notna()
            elif fill_missing_months and prof is not None:
                d = first_weekday(y, m, wd_idx)
                vals = add_variation(prof, [False] * len(prof), rel_g, f"{eui}|after|{d}") if natural else prof.copy()
                mask = pd.Series(False, index=prof.index)
                src = f"predicted - no data this month (typical {weekday})"
            else:
                continue
            start = b_date if b_meas is not None else d
            results.append(dict(
                before=(b_date, b_vals, b_mask) if b_meas is not None else None,
                name=f"result_{eui}_{start:%Y-%m-%d}_{d:%Y-%m-%d}.xlsx", day=d, eui=eui, month=f"{y}-{m:02d}",
                basis=basis, values=[None if pd.isna(v) else int(round(v)) for v in vals],
                measured=[bool(x) for x in mask], source=src))
    return results


def workbook_bytes(res, highlight=True, header_mode="date"):
    """Plain, portal-friendly sheet: Time | Before date | After date, 90 data rows only,
    left-aligned, no frozen panes (those show as thick lines), no extra formatting beyond the data."""
    wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
    d, bf = res["day"], res.get("before")
    # header_mode: "date" = real Excel date cells (B1, C1) | "weekday" = text "25/09/2025 (Thursday)" | "text" = text "25/09/2025"
    def lab(x):
        if header_mode == "date":
            return x if isinstance(x, date) else x.date()
        return f"{x:%d/%m/%Y} ({pd.Timestamp(x).day_name()})" if header_mode == "weekday" else f"{x:%d/%m/%Y}"
    head = ["Time"] + ([lab(bf[0])] if bf else []) + [lab(d)]
    left = Alignment(horizontal="left")
    ws.append(head)
    for i, t in enumerate(SLOTS):
        ws.append([t] + ([bf[1][i]] if bf else []) + [res["values"][i]])
        ws.cell(row=i + 2, column=1).number_format = "h:mm:ss"
        if highlight:
            if bf and not bf[2][i]:
                ws.cell(row=i + 2, column=2).fill = YELLOW
            if not res["measured"][i]:
                ws.cell(row=i + 2, column=len(head)).fill = YELLOW
    for row in ws.iter_rows(min_row=1, max_row=len(SLOTS) + 1, max_col=len(head)):
        for c in row:
            c.alignment = left
    if header_mode == "date":
        for col in range(2, len(head) + 1):
            ws.cell(row=1, column=col).number_format = "dd/mm/yyyy"
    ws.column_dimensions["A"].width = 10
    for col in range(2, len(head) + 1):
        ws.column_dimensions[chr(64 + col)].width = 14
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def log_frame(results):
    return pd.DataFrame([dict(month=r["month"], file=r["name"], date=f"{r['day']:%Y-%m-%d}",
                              weekday=WD[r["day"].weekday()], source=r["source"],
                              pattern_based_on=r["basis"], measured_slots=sum(r["measured"]), predicted_slots=len(SLOTS) - sum(r["measured"]),
                              before_predicted_slots=(len(SLOTS) - sum(r["before"][2])) if r.get("before") else "")
                         for r in results])


def build_zip(results, highlight=True, header_mode="date"):
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        for r in results:
            z.writestr(r["name"], workbook_bytes(r, highlight, header_mode))
        z.writestr("estimate_log.csv", log_frame(results).to_csv(index=False))
    return b.getvalue()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("csv"); ap.add_argument("-o", "--out", default="monthly_files")
    ap.add_argument("--weekday", default="Thu"); ap.add_argument("--min-coverage", type=float, default=0.2)
    ap.add_argument("--value", default="duration_in_traffic")
    ap.add_argument("--no-fill-months", action="store_true"); ap.add_argument("--no-highlight", action="store_true")
    ap.add_argument("--before", help="Before file (csv/xlsx with Time + '(Before)' column); added to every month")
    ap.add_argument("--header", choices=["date", "weekday", "text"], default="date",
                    help="date = real date cells (default) | weekday = text '25/09/2025 (Thursday)' | text = text '25/09/2025'")
    ap.add_argument("--no-variation", action="store_true", help="leave predicted cells as smooth values (no natural scatter)")
    ap.add_argument("--before-date", help="Before date as YYYY-MM-DD if it is not in the column name")
    ap.add_argument("--force-weekday", action="store_true",
                    help="if a month has no preferred-weekday data, predict that weekday instead of using another real day")
    a = ap.parse_args()
    bf = load_before(a.before, date.fromisoformat(a.before_date) if a.before_date else None) if a.before else None
    res = process(load(a.csv, a.value), a.weekday, a.min_coverage, not a.no_fill_months, a.force_weekday, bf, not a.no_variation)
    os.makedirs(a.out, exist_ok=True)
    for r in res:
        open(os.path.join(a.out, r["name"]), "wb").write(workbook_bytes(r, not a.no_highlight, a.header))
    lg = log_frame(res); lg.to_csv(os.path.join(a.out, "estimate_log.csv"), index=False)
    print(lg.to_string(index=False))


if __name__ == "__main__":
    sys.exit(main())
